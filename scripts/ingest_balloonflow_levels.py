#!/usr/bin/env python
"""
Ingest BalloonFlow level data (xlsx + LevelDB backups + PixelArtConverter cs)
into MongoDB.

Splits:
- BalloonFlow_LevelDesign.xlsx           -> design_base (data_type=table, system=LevelDesignSheet)
- LevelDB_backup_*.json (52 backups)     -> level_backups collection (1 entry per backup, schema summary)
- PixelArtConverter .cs files (4 files)  -> code_base (Generic genre, project=BalloonFlow_PixelArtConverter)
"""
import json
import re
import hashlib
from pathlib import Path
from datetime import datetime, timezone

from pymongo import MongoClient
import openpyxl

ENV_PATH = Path("E:/AI/.env")
env = {}
for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
    m = re.match(r"^([^=]+)=(.*)$", line)
    if m:
        env[m.group(1).strip()] = m.group(2).strip()

MONGO_URI = env["MONGO_URI"]
PROJECT = "BalloonFlow"
GENRE = "Puzzle"
NOW = datetime.now(timezone.utc).isoformat()

XLSX_PATH = Path("E:/Balloonflow/BalloonFlow_Level/BalloonFlow_LevelDesign.xlsx")
LEVEL_BACKUP_DIR = Path("E:/Balloonflow/BalloonFlow_History/LevelBackups_backup_20260407_105754")
CONVERTER_DIR = Path("E:/Balloonflow/BalloonFlow_Level/Converter/PixelArtConverter")


def hash12(s: str) -> str:
    return hashlib.sha1(s.encode("utf-8")).hexdigest()[:12]


def ingest_xlsx(db):
    if not XLSX_PATH.exists():
        print("SKIP: xlsx not found")
        return

    wb = openpyxl.load_workbook(str(XLSX_PATH), data_only=True, read_only=True)
    sheet_summaries = []
    full_text_parts = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        sheet_summaries.append({
            "sheet": sn,
            "row_count": len(rows),
            "first_rows": rows[:10],
        })
        full_text_parts.append(f"\n=== Sheet: {sn} ({len(rows)} rows) ===")
        full_text_parts.extend(rows[:200])  # cap per sheet
    wb.close()

    text = "\n".join(full_text_parts)
    doc = {
        "designId": "balloonflow_content_office_leveldesignsheet",
        "project": PROJECT,
        "genre": GENRE,
        "domain": "Content",
        "system": "LevelDesignSheet",
        "data_type": "table",
        "score": 0.5,
        "source": "internal_original",
        "source_file": str(XLSX_PATH),
        "document_role": "level_data_sheet",
        "tags": ["level_design", "spreadsheet", "300_levels"],
        "content": text[:200000],
        "content_hash": hash12(text),
        "summary": f"BalloonFlow level design sheet — {len(sheet_summaries)} sheets: {[s['sheet'] for s in sheet_summaries]}",
        "sheets": sheet_summaries,
        "char_count": len(text),
        "provides": [], "requires": [],
        "version": 1, "timestamp": NOW, "ingested_at": NOW,
    }
    coll = db["design_base"]
    result = coll.update_one(
        {"designId": doc["designId"], "project": PROJECT},
        {"$set": doc, "$unset": {"embedding": "", "embedded_at": "", "embedding_model": ""}},
        upsert=True,
    )
    print(f"  xlsx → design_base: {doc['designId']} ({'added' if result.upserted_id else 'updated'})")


def ingest_level_backups(db):
    coll = db["level_backups"]
    try: coll.create_index([("backupId", 1), ("project", 1)], unique=True)
    except Exception: pass
    try: coll.create_index([("project", 1), ("timestamp", -1)])
    except Exception: pass

    if not LEVEL_BACKUP_DIR.exists():
        print("SKIP: level backups dir not found")
        return

    files = sorted([p for p in LEVEL_BACKUP_DIR.iterdir() if p.suffix == ".json"])
    print(f"\n  Found {len(files)} level backups")
    added = 0
    for fp in files:
        try:
            text = fp.read_text(encoding="utf-8")
            data = json.loads(text)
        except Exception as e:
            print(f"    FAIL parse: {fp.name}: {e}")
            continue

        m = re.search(r"LevelDB_backup_(\d{8}_\d{6})\.json", fp.name)
        ts = m.group(1) if m else fp.stem
        backup_id = f"balloonflow_levelbackup_{ts}"

        # Schema discovery
        level_count = 0
        first_level_keys = []
        if isinstance(data, list):
            level_count = len(data)
            if data and isinstance(data[0], dict):
                first_level_keys = list(data[0].keys())[:20]
        elif isinstance(data, dict):
            for k, v in data.items():
                if isinstance(v, list) and len(v) > level_count:
                    level_count = len(v)
                    if v and isinstance(v[0], dict):
                        first_level_keys = list(v[0].keys())[:20]

        doc = {
            "backupId": backup_id,
            "project": PROJECT,
            "genre": GENRE,
            "timestamp": ts,
            "file_path": str(fp).replace("\\", "/"),
            "file_size_bytes": fp.stat().st_size,
            "level_count": level_count,
            "schema_keys": first_level_keys,
            "ingested_at": NOW,
        }
        result = coll.update_one(
            {"backupId": backup_id, "project": PROJECT},
            {"$set": doc},
            upsert=True,
        )
        if result.upserted_id: added += 1
    print(f"  level_backups: {added} added (total scanned {len(files)})")


def ingest_converter_cs(db):
    coll = db["code_base"]
    if not CONVERTER_DIR.exists():
        print("SKIP: converter dir not found")
        return

    cs_files = list(CONVERTER_DIR.glob("*.cs"))
    print(f"\n  Found {len(cs_files)} converter .cs files")
    added = updated = 0
    for fp in cs_files:
        text = fp.read_text(encoding="utf-8", errors="ignore")
        cls_match = re.search(r"\b(?:public|internal)?\s*(?:partial\s+)?(?:abstract\s+)?class\s+(\w+)", text)
        class_name = cls_match.group(1) if cls_match else fp.stem
        provides = [m.group(1) for m in re.finditer(r"\bpublic\s+(?:static\s+|async\s+|virtual\s+|override\s+)*\w[\w<>,\s\[\]]*?\s+(\w+)\s*\(", text)][:20]

        file_id = hash12(f"BalloonFlow_PixelArtConverter:{fp.name}")
        rel = str(fp).replace("\\", "/")

        doc = {
            "fileId": file_id,
            "project": "BalloonFlow_PixelArtConverter",
            "genre": "Generic",  # WPF tool, not Unity
            "layer": "Game",
            "role": "Tool",
            "system": "PixelArtConverter",
            "className": class_name,
            "file_name": fp.name,
            "file_path": rel,
            "score": 0.5,
            "summary": f"BalloonFlow PixelArtConverter tool — {class_name} (WPF level data exporter)",
            "tags": ["tool", "level_converter", "pixel_art", "wpf"],
            "contract": {"provides": provides, "requires": []},
            "content": text[:50000],
            "content_hash": hash12(text),
            "char_count": len(text),
            "timestamp": NOW,
            "ingested_at": NOW,
        }
        result = coll.update_one(
            {"fileId": file_id, "project": doc["project"]},
            {"$set": doc, "$unset": {"embedding": "", "embedded_at": "", "embedding_model": ""}},
            upsert=True,
        )
        if result.upserted_id: added += 1
        elif result.modified_count: updated += 1
    print(f"  converter cs: added={added}, updated={updated}")


def main():
    client = MongoClient(MONGO_URI)
    db = client["aigame"]
    print("=== xlsx LevelDesign ===")
    ingest_xlsx(db)
    print("=== Level backups ===")
    ingest_level_backups(db)
    print("=== PixelArtConverter cs ===")
    ingest_converter_cs(db)
    client.close()


if __name__ == "__main__":
    main()
