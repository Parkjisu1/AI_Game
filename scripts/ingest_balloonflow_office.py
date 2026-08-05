#!/usr/bin/env python
"""
Ingest BalloonFlow .docx/.xlsx files into design_base.
Extracts text from docx/xlsx and upserts as design_base entries.
"""
import sys
import os
import hashlib
import re
import io
from pathlib import Path
from datetime import datetime, timezone

# pymongo + python-docx + openpyxl
from pymongo import MongoClient
from docx import Document
import openpyxl

ENV_PATH = Path("E:/AI/.env")
env = {}
for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
    m = re.match(r"^([^=]+)=(.*)$", line)
    if m:
        env[m.group(1).strip()] = m.group(2).strip()

MONGO_URI = env["MONGO_URI"]
DOCS_DIRS = [
    Path("E:/Balloonflow/BalloonFlow_Doc/Recent"),
    Path("E:/Balloonflow/BalloonFlow_Doc/2026-05-28"),
]
PROJECT = "BalloonFlow"
GENRE = "Puzzle"
NOW = datetime.now(timezone.utc).isoformat()

FILE_MAP = {
    "BalloonFlow_컨셉문서.docx": {
        "domain": "InGame",
        "system": "ConceptDocFull",
        "role": "concept",
        "tags": ["concept", "core_pillar", "vision"],
    },
    "퍼즐_장르_설계표준.docx": {
        "domain": "Meta",
        "system": "PuzzleGenreStandard",
        "role": "standard",
        "tags": ["genre_standard", "design_methodology", "puzzle"],
    },
    "balloonflow_IAP.xlsx": {
        "domain": "BM",
        "system": "IAPCatalog",
        "role": "data_table",
        "tags": ["bm", "iap", "monetization", "price_table"],
    },
    # --- 2026-06-02 신규 (레벨 데이터 catalog은 제외; QA 문서만 포함) ---
    "00-BalloonFlow_QA_테스트케이스 - 1.0.xlsx": {
        "domain": "Meta",
        "system": "QATestCases",
        "role": "test_cases",
        "tags": ["qa", "testing", "test_cases", "checklist"],
    },
}


def hash_id(s: str) -> str:
    return hashlib.sha1(s.encode("utf-8")).hexdigest()[:12]


def extract_docx(path: Path) -> tuple[str, list[dict]]:
    doc = Document(str(path))
    paragraphs = []
    headings = []
    for p in doc.paragraphs:
        txt = (p.text or "").strip()
        if not txt:
            continue
        style = (p.style.name or "") if p.style else ""
        if style.lower().startswith("heading"):
            level = 1
            m = re.search(r"(\d)", style)
            if m:
                level = int(m.group(1))
            headings.append({"level": level, "title": txt})
        paragraphs.append(txt)
    # Tables
    for ti, tbl in enumerate(doc.tables):
        paragraphs.append(f"\n--- Table {ti+1} ---")
        for row in tbl.rows:
            cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
            paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs), headings


def extract_xlsx(path: Path) -> tuple[str, list[dict]]:
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = []
    headings = []
    for sheet_name in wb.sheetnames:
        headings.append({"level": 1, "title": sheet_name})
        out.append(f"\n=== Sheet: {sheet_name} ===")
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out), headings


def make_summary(text: str, max_len: int = 600) -> str:
    paragraphs = [p.strip() for p in text.split("\n") if p.strip() and not p.strip().startswith("---")]
    return "\n".join(paragraphs[:5])[:max_len]


def main():
    client = MongoClient(MONGO_URI)
    db = client["aigame"]
    coll = db["design_base"]

    added = updated = 0
    for fname, meta in FILE_MAP.items():
        fpath = next((d / fname for d in DOCS_DIRS if (d / fname).exists()), None)
        if fpath is None:
            print(f"  SKIP (missing): {fname}")
            continue

        if fpath.suffix == ".docx":
            text, headings = extract_docx(fpath)
            data_type = "concept" if meta["role"] == "concept" else "standard"
        elif fpath.suffix == ".xlsx":
            text, headings = extract_xlsx(fpath)
            data_type = "table"
        else:
            continue

        sys_id = re.sub(r"[^a-z0-9]+", "_", meta["system"].lower()).strip("_")
        design_id = f"balloonflow_{meta['domain'].lower()}_office_{sys_id}"

        doc = {
            "designId": design_id,
            "project": PROJECT,
            "genre": GENRE,
            "domain": meta["domain"],
            "system": meta["system"],
            "data_type": data_type,
            "score": 0.5,
            "source": "internal_original",
            "source_file": str(fpath),
            "document_role": meta["role"],
            "tags": meta["tags"],
            "content": text[:200000],  # cap raw content
            "content_hash": hash_id(text),
            "headings": headings[:50],
            "summary": make_summary(text),
            "char_count": len(text),
            "provides": [],
            "requires": [],
            "version": 1,
            "timestamp": NOW,
            "ingested_at": NOW,
        }

        result = coll.update_one(
            {"designId": design_id, "project": PROJECT},
            {"$set": doc, "$unset": {"embedding": "", "embedded_at": "", "embedding_model": ""}},
            upsert=True,
        )
        if result.upserted_id:
            print(f"  + {design_id}  ({len(text)} chars)")
            added += 1
        elif result.modified_count:
            print(f"  ~ {design_id}  ({len(text)} chars)")
            updated += 1

    print(f"\nDone: added={added}, updated={updated}")
    client.close()


if __name__ == "__main__":
    main()
